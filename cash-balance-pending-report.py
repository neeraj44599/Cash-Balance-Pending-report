from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError
from datetime import datetime
import os
import time
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import requests
import matplotlib.pyplot as plt

# ============================================================
# CONFIGURATION
# ============================================================
USERNAME = os.getenv("ACEN_USERNAME")
PASSWORD = os.getenv("ACEN_PASSWORD")
LOGIN_URL = os.getenv("ACEN_LOGIN_URL")

TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
TELEGRAM_CHAT_ID = os.getenv("TELEGRAM_CHAT_ID")

BASE_DIR = os.getcwd()

HR_FILE_PATH = os.path.join(
    BASE_DIR,
    "Weekly HR Update as on 17062026.xlsx.xlsx"
)

# आपकी पसंद के अनुसार फाइनल पेंडिंग एक्सेल रिपोर्ट का नाम
OUTPUT_FILE_PATH = os.path.join(BASE_DIR, "Cash_Balance_Pending_Report.xlsx")
SCREENSHOT_PATH = os.path.join(BASE_DIR, "pending_report_screenshot.png")

# ============================================================
# VALIDATE CONFIGURATION
# ============================================================
def validate_configuration():
    required = {
        "ACEN_USERNAME": USERNAME,
        "ACEN_PASSWORD": PASSWORD,
        "ACEN_LOGIN_URL": LOGIN_URL,
        "TELEGRAM_BOT_TOKEN": TELEGRAM_BOT_TOKEN,
        "TELEGRAM_CHAT_ID": TELEGRAM_CHAT_ID
    }
    missing = [key for key, value in required.items() if not value]
    if missing:
        raise Exception("Missing GitHub Secrets: " + ", ".join(missing))
    print("All GitHub Secrets are available.")

# ============================================================
# EMPLOYEE ID NORMALIZATION
# ============================================================
def normalize_employee_id(value):
    if pd.isna(value):
        return ""
    value = str(value).strip().upper()
    value = value.replace("\u00A0", "").replace(" ", "").replace("\t", "").replace("\n", "").replace("\r", "")
    if value.endswith(".0"):
        value = value[:-2]
    return value

# ============================================================
# STATUS CLEANING
# ============================================================
def clean_status(value):
    if pd.isna(value):
        return ""
    return str(value).strip().replace("\u00A0", " ")

# ============================================================
# TELEGRAM FUNCTIONS
# ============================================================
def telegram_request(url, data=None, files=None):
    try:
        response = requests.post(url, data=data, files=files, timeout=60)
        if not response.ok:
            print("Telegram API Error:", response.status_code, response.text)
            return False
        return True
    except Exception as e:
        print("Telegram Request Exception:", e)
        return False

def send_telegram_report(bot_token, chat_id, message, excel_path=None, screenshot_path=None):
    if not bot_token or not chat_id:
        print("Telegram credentials missing. Telegram dispatch skipped.")
        return

    base_url = f"https://api.telegram.org/bot{bot_token}"

    # 1. SEND SCREENSHOT WITH CAPTION
    if screenshot_path and os.path.exists(screenshot_path):
        print("Sending Excel screenshot with report caption...")
        try:
            with open(screenshot_path, "rb") as photo:
                telegram_request(
                    f"{base_url}/sendPhoto",
                    data={"chat_id": chat_id, "caption": message, "parse_mode": "Markdown"},
                    files={"photo": photo}
                )
        except Exception as e:
            print("Screenshot Telegram Error:", e)

    # 2. SEND EXCEL FILE
    if excel_path and os.path.exists(excel_path):
        print("Sending Excel file...")
        try:
            with open(excel_path, "rb") as document:
                telegram_request(
                    f"{base_url}/sendDocument",
                    data={"chat_id": chat_id},
                    files={"document": document}
                )
        except Exception as e:
            print("Excel Telegram Error:", e)

    print("Telegram dispatch completed.")

# ============================================================
# CREATE EXCEL SCREENSHOT (YOUR MATPLOTLIB LOGIC)
# ============================================================
def create_excel_screenshot(df, screenshot_path):
    print("Creating Excel-style screenshot...")
    if df.empty:
        print("No records available for screenshot.")
        return

    display_df = df.copy()
    preferred_columns = [
        "State", "Division", "Cluster", "Branch", "Branch Code", 
        "Employee", "Employee ID", "Designation", "Opening Balance", 
        "Collection", "Deposit", "Closing Balance", "Status"
    ]
    screenshot_columns = [col for col in preferred_columns if col in display_df.columns]
    
    # अधिकतम टॉप 15 रिकॉर्ड्स दिखाना स्क्रीनशॉट इमेज के लिए
    display_df = display_df[screenshot_columns].head(15)

    fig, ax = plt.subplots(figsize=(14, len(display_df) * 0.4 + 1.5))
    ax.axis('tight')
    ax.axis('off')

    table = ax.table(
        cellText=display_df.values,
        colLabels=display_df.columns,
        cellLoc='center',
        loc='center'
    )

    # स्टाइलिंग
    table.auto_set_font_size(False)
    table.set_font_size(9)
    table.scale(1.2, 1.3)

    for (row, col), cell in table.get_celld().items():
        if row == 0:
            cell.set_text_props(weight='bold', color='white')
            cell.set_facecolor('#4F81BD')
        else:
            # स्टेटस 'Inactive' होने पर ग्रे कलर बैकग्राउंड सेट करना
            status_val = str(display_df.iloc[row-1]['Status']).strip().lower() if 'Status' in display_df.columns else ''
            if 'inactive' in status_val or 'grey' in status_val:
                cell.set_facecolor('#D9D9D9') # Grey background for Inactive
            elif row % 2 == 0:
                cell.set_facecolor('#F2F2F2')
            else:
                cell.set_facecolor('white')

    plt.savefig(screenshot_path, bbox_inches='tight', dpi=200)
    plt.close()
    print("Screenshot created successfully.")

# ============================================================
# MAIN AUTOMATION & FILTER PROCESS
# ============================================================
def run_cash_balance_report_job():
    validate_configuration()
    
    # अस्थाई रॉ डाउनलोड पाथ
    raw_download_path = os.path.join(BASE_DIR, "temp_raw_data.xlsx")

    with sync_playwright() as p:
        print("Browser ओपन किया जा रहा है...")
        browser = p.chromium.launch(headless=True, args=["--no-sandbox"])
        context = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        )
        page = context.new_page()
        
        print(f"लॉगिन पेज ओपन किया जा रहा है: {LOGIN_URL}")
        page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=120000)
        time.sleep(5)
        
        target = page
        if len(page.frames) > 1:
            for f in page.frames:
                if f.name or f.url:
                    target = f
                    break

        try:
            print("लॉगिन डिटेल्स भरी जा रही हैं...")
            target.locator("input[type='text'], input[type='email'], input[name='username'], #username").first.fill(str(USERNAME))
            target.locator("input[type='password'], #password").first.fill(str(PASSWORD))
            time.sleep(1)
            target.locator("button[type='submit'], input[type='submit'], button:has-text('Login')").first.click()
            print("लॉगिन सबमिट हो गया।")
            time.sleep(10)
        except Exception as e:
            print("लॉगिन स्टेप फेल या स्किप:", e)

        for f in page.frames:
            if "main" in f.name.lower() or "report" in f.name.lower() or len(f.name) > 0:
                target = f
                break

        print("Reports बटन पर जाया जा रहा है...")
        target.get_by_text("Reports", exact=True).wait_for(state="visible", timeout=60000)
        target.get_by_text("Reports", exact=True).click()
        time.sleep(4)

        print("डाउनलोड बटन पर क्लिक किया जा रहा है...")
        try:
            download_btn = target.locator("button, a").filter(has_text=re.compile(r"download|export|excel|एक्सेल", re.IGNORECASE)).first
            with page.expect_download(timeout=60000) as download_info:
                download_btn.click()
            download = download_info.value
            download.save_as(raw_download_path)
            print("फ़ाइल डाउनलोड हो गई।")
        except Exception as e:
            print("डाउनलोड फ़ेल, अल्टरनेटिव तरीका आज़माया जा रहा है...")
            with page.expect_download(timeout=30000) as download_info:
                page.locator("a:has-text('Download'), button:has-text('Download')").first.click()
            download = download_info.value
            download.save_as(raw_download_path)
            
        browser.close()

    # ============================================================
    # PANDAS DATA PROCESSING & FILTERING (A to Z FIXED)
    # ============================================================
    if not os.path.exists(raw_download_path):
        print("Error: एक्सेल फ़ाइल नहीं मिली! डाउनलोड फेल हुआ।")
        return

    print("डाउनलोड की गई फ़ाइल को फ़िल्टर और सॉर्ट किया जा रहा है...")
    df = pd.read_excel(raw_download_path)
    df.columns = df.columns.str.strip()
    
    # कॉलम नाम ऑटो-डिटेक्ट करना (Pending या Cash, और Employee या Name के लिए)
    pending_col = next((c for c in df.columns if 'pending' in c.lower() or 'cash' in c.lower() or 'closing' in c.lower()), 'Closing Balance')
    emp_col = next((c for c in df.columns if 'emp' in c.lower() or 'name' in c.lower()), 'Employee')
    status_col = next((c for c in df.columns if 'status' in c.lower() or 'state' in c.lower()), 'Status')
    
    df[pending_col] = pd.to_numeric(df[pending_col], errors='coerce').fillna(0)
    df[emp_col] = df[emp_col].astype(str).str.strip()
    if status_col in df.columns:
    df[status_col] = df[status_col].apply(clean_status)
    # 💡 फ़िल्टर: सिर्फ Pending 0 वालों को हटाना है (Active और Inactive Grey वाले दोनों रहेंगे)
    filtered_df = df[df[pending_col] > 0].copy()
    # 💡 ऑटो सॉर्टिंग: पैसों के बजाय कर्मचारियों के नाम के हिसाब से A to Z क्रम में सेट करना
    filtered_df = filtered_df.sort_values(by=emp_col, ascending=True)
    # फाइनल एक्सेल फ़ाइल सेव करना
    filtered_df.to_excel(OUTPUT_FILE_PATH, index=False)
    print(f"नई फ़िल्टर्ड रिपोर्ट सेव हो गई: {OUTPUT_FILE_PATH}")
    # स्क्रीनशॉट इमेज बनाना
    create_excel_screenshot(filtered_df, SCREENSHOT_PATH)
    if os.path.exists(raw_download_path):
    os.remove(raw_download_path)
    total_cases = len(filtered_df)
    total_amount = filtered_df[pending_col].sum()
    summary_message = (
    "🌅 Cash Balance Pending Report (A to Z)\n"
    "━━━━━━━━━━━━━━━━━━━━━━━━\n"
    f"👥 Total Pending Employees: {total_cases}\n"
    f"💰 Total Pending Amount: ₹{total_amount:,.2f}\n"
    "━━━━━━━━━━━━━━━━━━━━━━━━\n"
    "📢 Note: इस लिस्ट में Active और Inactive दोनों कर्मचारी शामिल हैं। 0 पेंडिंग वाले लोगों को हटा दिया गया है और डेटा को नाम के अनुसार A to Z व्यवस्थित किया गया है।"
    )
    # टेलीग्राम पर पब्लिश करना (मेसेज + इमेज स्क्रीनशॉट + एक्सेल शीट)
    send_telegram_report(TELEGRAM_BOT_TOKEN, TELEGRAM_CHAT_ID, summary_message, OUTPUT_FILE_PATH, SCREENSHOT_PATH)
    if name == "main":
    try:
    run_cash_balance_report_job()
    except Exception as e:
    print(f"इस रिपोर्ट जॉब में एरर आया: {e}")
